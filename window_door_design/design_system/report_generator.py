import os
from io import BytesIO
from datetime import datetime
from django.conf import settings
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from design_system.models import WindowDoorDesign, DesignReport
from design_system.business_logic import DesignCalculator


class ReportGenerator:
    """Generates various reports (PDF, Excel) for window/door designs"""
    
    def __init__(self, design: WindowDoorDesign):
        self.design = design
        self.calculator = DesignCalculator(design)
    
    def generate_pdf_quotation(self):
        """Generate quotation as PDF"""
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter)
        elements = []
        
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=24,
            textColor=colors.HexColor('#1f4788'),
            spaceAfter=30,
            alignment=1  # Center
        )
        
        heading_style = ParagraphStyle(
            'CustomHeading',
            parent=styles['Heading2'],
            fontSize=14,
            textColor=colors.HexColor('#333333'),
            spaceAfter=12,
            spaceBefore=12
        )
        
        # Title
        elements.append(Paragraph("QUOTATION", title_style))
        elements.append(Spacer(1, 0.2 * 72))
        
        # Header info
        quotation_data = self.calculator.get_quotation_data()
        header_data = [
            ['Design Name:', quotation_data['design_name']],
            ['Design Code:', quotation_data['design_code']],
            ['Type:', quotation_data['description']],
            ['Dimensions:', quotation_data['dimensions']],
            ['Date:', datetime.now().strftime('%Y-%m-%d')],
        ]
        
        header_table = Table(header_data, colWidths=[2 * 72, 4 * 72])
        header_table.setStyle(TableStyle([
            ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
            ('TOPPADDING', (0, 0), (-1, -1), 8),
        ]))
        
        elements.append(header_table)
        elements.append(Spacer(1, 0.3 * 72))
        
        # Cost breakdown
        elements.append(Paragraph("Cost Breakdown", heading_style))
        
        cost_data = [
            ['Item', 'Amount'],
            ['Material Cost', f"${quotation_data['cost_breakdown']['material_cost']}"],
            ['Labor Cost', f"${quotation_data['cost_breakdown']['labor_cost']}"],
            ['Glass Cost', f"${quotation_data['cost_breakdown']['glass_cost']}"],
            ['Mesh Cost', f"${quotation_data['cost_breakdown']['mesh_cost']}"],
            ['Profile Cost', f"${quotation_data['cost_breakdown']['profile_cost']}"],
            ['TOTAL', f"${quotation_data['cost_breakdown']['total_cost']}"],
        ]
        
        cost_table = Table(cost_data, colWidths=[3 * 72, 2 * 72])
        cost_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1f4788')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'RIGHT'),
            ('ALIGN', (0, 0), (0, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 11),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 10),
            ('TOPPADDING', (0, 0), (-1, -1), 10),
            ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#e8eef7')),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
        ]))
        
        elements.append(cost_table)
        
        doc.build(elements)
        buffer.seek(0)
        return buffer
    
    def generate_excel_boq(self):
        """Generate Bill of Quantities as Excel"""
        wb = Workbook()
        ws = wb.active
        ws.title = "BOQ"
        
        # Styles
        header_fill = PatternFill(start_color="1f4788", end_color="1f4788", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=12)
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Title
        ws['A1'] = "BILL OF QUANTITIES"
        ws['A1'].font = Font(bold=True, size=14, color="1f4788")
        ws.merge_cells('A1:D1')
        
        # Design details
        row = 3
        boq = self.calculator.get_bill_of_quantities()
        
        details = [
            ['Design Name:', boq['design_name']],
            ['Design Code:', boq['design_code']],
            ['Type:', boq['design_type']],
            ['Material:', boq['material_type']],
            ['Number of Units:', boq['number_of_units']],
            ['Dimensions:', f"{boq['width_m']}m x {boq['height_m']}m"],
            ['Total Area:', f"{boq['total_area_sqm']} sq.m"],
            ['Glass Type:', boq['glass_type']],
            ['Has Mesh:', 'Yes' if boq['has_mesh'] else 'No'],
            ['Finish Type:', boq['finish_type']],
        ]
        
        for label, value in details:
            ws[f'A{row}'] = label
            ws[f'B{row}'] = value
            ws[f'A{row}'].font = Font(bold=True)
            row += 1
        
        # Materials section
        row += 2
        ws[f'A{row}'] = "MATERIALS REQUIRED"
        ws[f'A{row}'].font = Font(bold=True, size=11, color="FFFFFF")
        ws[f'A{row}'].fill = header_fill
        ws.merge_cells(f'A{row}:D{row}')
        row += 1
        
        materials = boq['materials']
        materials_data = [
            ['Profile Length (m):', materials['profile_length_meters']],
            ['Glass Area (sq.m):', materials['glass_area_sqm']],
            ['Mesh Area (sq.m):', materials['mesh_area_sqm']],
            ['Frame Width (mm):', materials['frame_width_mm']],
            ['Frame Height (mm):', materials['frame_height_mm']],
        ]
        
        for label, value in materials_data:
            ws[f'A{row}'] = label
            ws[f'B{row}'] = value
            ws[f'A{row}'].font = Font(bold=True)
            row += 1
        
        # Cost breakdown section
        row += 2
        ws[f'A{row}'] = "COST BREAKDOWN"
        ws[f'A{row}'].font = Font(bold=True, size=11, color="FFFFFF")
        ws[f'A{row}'].fill = header_fill
        ws.merge_cells(f'A{row}:D{row}')
        row += 1
        
        costs = boq['costs']
        cost_data = [
            ['Material Cost:', f"${costs['material_cost']}"],
            ['Labor Cost:', f"${costs['labor_cost']}"],
            ['Glass Cost:', f"${costs['glass_cost']}"],
            ['Mesh Cost:', f"${costs['mesh_cost']}"],
            ['Profile Cost:', f"${costs['profile_cost']}"],
            ['TOTAL COST:', f"${costs['total_cost']}"],
        ]
        
        for label, value in cost_data:
            ws[f'A{row}'] = label
            ws[f'B{row}'] = value
            ws[f'A{row}'].font = Font(bold=True)
            if 'TOTAL' in label:
                ws[f'A{row}'].fill = PatternFill(start_color="e8eef7", end_color="e8eef7", fill_type="solid")
                ws[f'B{row}'].fill = PatternFill(start_color="e8eef7", end_color="e8eef7", fill_type="solid")
            row += 1
        
        # Adjust column widths
        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 25
        ws.column_dimensions['C'].width = 25
        ws.column_dimensions['D'].width = 25
        
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer
    
    def generate_excel_cutting_list(self):
        """Generate cutting list as Excel"""
        wb = Workbook()
        ws = wb.active
        ws.title = "Cutting List"
        
        header_fill = PatternFill(start_color="1f4788", end_color="1f4788", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=11)
        
        # Title
        ws['A1'] = "CUTTING LIST"
        ws['A1'].font = Font(bold=True, size=14, color="1f4788")
        ws.merge_cells('A1:E1')
        
        # Design info
        ws['A2'] = f"Design: {self.design.name} ({self.design.code})"
        ws['A3'] = f"Date: {datetime.now().strftime('%Y-%m-%d')}"
        
        # Cutting list header
        headers = ['Unit #', 'Vertical Rail (mm)', 'Horizontal Rail (mm)', 'Glass Width (mm)', 'Glass Height (mm)']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=5, column=col)
            cell.value = header
            cell.font = header_font
            cell.fill = header_fill
        
        # Cutting sizes
        cutting_sizes = self.calculator.calculate_cutting_sizes()
        for row_idx, cut in enumerate(cutting_sizes, 6):
            ws[f'A{row_idx}'] = cut['unit']
            ws[f'B{row_idx}'] = cut['vertical_rail_mm']
            ws[f'C{row_idx}'] = cut['horizontal_rail_mm']
            ws[f'D{row_idx}'] = cut['glass_width_mm']
            ws[f'E{row_idx}'] = cut['glass_height_mm']
        
        # Adjust column widths
        for col in ['A', 'B', 'C', 'D', 'E']:
            ws.column_dimensions[col].width = 20
        
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer
    
    def generate_material_summary_pdf(self):
        """Generate material summary report as PDF"""
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4)
        elements = []
        
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=20,
            textColor=colors.HexColor('#1f4788'),
            spaceAfter=20,
            alignment=1
        )
        
        # Title
        elements.append(Paragraph("MATERIAL SUMMARY REPORT", title_style))
        elements.append(Spacer(1, 0.2 * 72))
        
        # Summary data
        materials_req = self.calculator.calculate_material_requirements()
        
        summary_data = [
            ['Parameter', 'Value', 'Unit'],
            ['Design Name', self.design.name, ''],
            ['Design Code', self.design.code, ''],
            ['Type', self.design.get_type_choice_display(), ''],
            ['Material', self.design.get_material_display(), ''],
            ['Profile Length Required', materials_req['profile_length_meters'], 'meters'],
            ['Glass Area', materials_req['glass_area_sqm'], 'sq.m'],
            ['Mesh Area', materials_req['mesh_area_sqm'], 'sq.m'],
            ['Frame Width', materials_req['frame_width_mm'], 'mm'],
            ['Frame Height', materials_req['frame_height_mm'], 'mm'],
        ]
        
        summary_table = Table(summary_data, colWidths=[3 * 72, 2 * 72, 1 * 72])
        summary_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1f4788')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
            ('TOPPADDING', (0, 0), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
        ]))
        
        elements.append(summary_table)
        elements.append(Spacer(1, 0.3 * 72))
        
        doc.build(elements)
        buffer.seek(0)
        return buffer
